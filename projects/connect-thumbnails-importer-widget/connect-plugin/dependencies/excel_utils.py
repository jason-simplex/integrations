# :coding: utf-8
# :copyright: Copyright (c) 2025 bro.tiger
import logging
import os
import tempfile
import uuid
import zipfile
import posixpath
import xml.etree.ElementTree as ET
from typing import Dict, List, Optional, Tuple

logger = logging.getLogger('batch_import_thumbnails')



def extract_original_images_by_row_from_xlsx(excel_path: str, sheet_name: str,
                                             image_col_name: str, headers: List[str]) -> Dict[int, str]:
    """Extract original embedded images mapped by row for a given sheet and image column.

    Parses the .xlsx structure to find drawing anchors and resolves image relationship targets,
    preserving original resolution and format. Only accepts twoCellAnchor staying within the
    same cell (from/to same column and row).
    """
    mapping: Dict[int, str] = {}
    try:
        z = zipfile.ZipFile(excel_path)
    except Exception as e:
        logger.warning(f'Failed to open xlsx as zip: {e}')
        return mapping

    ns_main = {'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
               'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    ns_xdr = {'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
              'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
              'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}

    try:
        workbook_xml = z.read('xl/workbook.xml')
        workbook = ET.fromstring(workbook_xml)
        rels_xml = z.read('xl/_rels/workbook.xml.rels')
        rels = ET.fromstring(rels_xml)
    except Exception as e:
        logger.warning(f'Failed to read workbook relations: {e}')
        return mapping

    rid_to_target = {}
    for rel in rels.findall('.//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship'):
        rid = rel.get('Id')
        target = rel.get('Target')
        if rid and target:
            rid_to_target[rid] = target  # e.g. 'worksheets/sheet1.xml'

    sheet_xml_path = None
    for sheet in workbook.findall('.//main:sheets/main:sheet', ns_main):
        name = sheet.get('name')
        rid = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        if name == sheet_name and rid in rid_to_target:
            sheet_xml_path = 'xl/' + rid_to_target[rid]
            break

    if not sheet_xml_path:
        logger.warning(f'Sheet "{sheet_name}" not found in workbook rels.')
        return mapping

    try:
        sheet_xml = ET.fromstring(z.read(sheet_xml_path))
    except Exception as e:
        logger.warning(f'Failed to read sheet xml: {e}')
        return mapping

    drawing_elem = sheet_xml.find('.//main:drawing', ns_main)
    if drawing_elem is None:
        return mapping
    drawing_rid = drawing_elem.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
    sheet_rels_path = sheet_xml_path.replace('worksheets/', 'worksheets/_rels/').replace('.xml', '.xml.rels')
    try:
        sheet_rels_xml = ET.fromstring(z.read(sheet_rels_path))
    except Exception as e:
        logger.warning(f'Failed to read sheet rels: {e}')
        return mapping
    drawing_target = None
    for rel in sheet_rels_xml.findall('.//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship'):
        if rel.get('Id') == drawing_rid:
            drawing_target = rel.get('Target')
            break
    if not drawing_target:
        return mapping
    # Resolve drawing target relative to the sheet part directory
    drawing_xml_path = posixpath.normpath(
        posixpath.join(posixpath.dirname(sheet_xml_path), drawing_target)
    )

    try:
        drawing_xml = ET.fromstring(z.read(drawing_xml_path))
        drawing_rels_path = drawing_xml_path.replace('drawings/', 'drawings/_rels/').replace('.xml', '.xml.rels')
        drawing_rels_xml = ET.fromstring(z.read(drawing_rels_path))
    except Exception as e:
        logger.warning(f'Failed to read drawing xml/relations: {e}')
        return mapping

    drid_to_target = {}
    for rel in drawing_rels_xml.findall('.//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship'):
        rid = rel.get('Id')
        target = rel.get('Target')
        if rid and target:
            # Resolve relative to the drawing part directory
            full = posixpath.normpath(
                posixpath.join(posixpath.dirname(drawing_xml_path), target)
            )
            drid_to_target[rid] = full

    try:
        image_col_zero = headers.index(image_col_name)
    except ValueError:
        logger.warning(f'Header "{image_col_name}" not found among {headers}.')
        return mapping

    tmpdir = tempfile.mkdtemp(prefix='excel_imgs_')

    def handle_anchor(anchor):
        frm = anchor.find('xdr:from', ns_xdr)
        if frm is None:
            return
        try:
            col = int(frm.find('xdr:col', ns_xdr).text)
            row = int(frm.find('xdr:row', ns_xdr).text)
        except Exception:
            return
        if col != image_col_zero:
            return
        # Ensure two-cell anchors stay within the same cell
        to = anchor.find('xdr:to', ns_xdr)
        if to is not None:
            try:
                to_col = int(to.find('xdr:col', ns_xdr).text)
                to_row = int(to.find('xdr:row', ns_xdr).text)
            except Exception:
                to_col = None
                to_row = None
            if to_col is None or to_row is None or to_col != image_col_zero or to_row != row:
                return
        pic = anchor.find('xdr:pic', ns_xdr)
        if pic is None:
            return
        blip = pic.find('xdr:blipFill', ns_xdr)
        if blip is None:
            return
        a_blip = blip.find('a:blip', ns_xdr)
        if a_blip is None:
            return
        rid = a_blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        if not rid or rid not in drid_to_target:
            return
        media_path = drid_to_target[rid]
        try:
            data = z.read(media_path)
        except Exception:
            return
        ext = os.path.splitext(media_path)[1] or '.png'
        out_path = os.path.join(tmpdir, f'{uuid.uuid4().hex}{ext}')
        try:
            with open(out_path, 'wb') as f:
                f.write(data)
        except Exception:
            return
        mapping[row + 1] = out_path  # XML rows are 0-based

    # Only consider twoCellAnchor for strict same-cell mapping
    for anc in drawing_xml.findall('.//xdr:twoCellAnchor', ns_xdr):
        handle_anchor(anc)

    return mapping

def read_headers_all(path: str):
    sheets: List[str] = []
    headers_map: Dict[str, List[str]] = {}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        sheets = list(getattr(wb, 'sheetnames', []))
        for name in sheets:
            try:
                ws = wb[name]
                it = ws.iter_rows(min_row=1, max_row=1, values_only=True)
                first = next(it, None)
                if first is not None:
                    headers = [str(v).strip() if v is not None else '' for v in (list(first) if isinstance(first, (list, tuple)) else [first])]
                else:
                    headers = []
                headers_map[name] = headers
            except Exception:
                headers_map[name] = []
        try:
            wb.close()
        except Exception:
            pass
    except Exception as e:
        logger.warning(f'read_headers_all failed: {e}')
    return sheets, headers_map
